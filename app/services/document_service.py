from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterator
from datetime import UTC
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from app.models import Registration


class DocumentTemplateError(Exception):
    """Raised when a registration document template cannot be used safely."""


TEMPLATE_ERROR_MESSAGE = "Document template could not be processed"
PLACEHOLDER_PATTERN = re.compile(r"{{[^{}]+}}")


def iter_paragraphs(parent: DocumentObject | _Cell) -> Iterator[Paragraph]:
    yield from parent.paragraphs
    for table in parent.tables:
        yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from iter_paragraphs(cell)


def replacement_values(
    registration: Registration, verified_by_username: str
) -> dict[str, str]:
    if registration.verified_at is None:
        raise DocumentTemplateError("Document template could not be processed")
    verified_at = registration.verified_at
    if verified_at.tzinfo is None:
        verified_at = verified_at.replace(tzinfo=UTC)
    return {
        "{{REGISTRATION_CODE}}": registration.registration_code,
        "{{FIRST_NAME}}": registration.first_name,
        "{{LAST_NAME}}": registration.last_name,
        "{{DATE_OF_BIRTH}}": registration.date_of_birth.isoformat(),
        "{{NATIONALITY}}": registration.nationality,
        "{{PASSPORT_NUMBER}}": registration.passport_number,
        "{{PHONE_NUMBER}}": registration.phone_number,
        "{{PARENT_PHONE_NUMBER}}": registration.parent_phone_number,
        "{{HSK_LEVEL}}": registration.hsk_level.value,
        "{{VERIFIED_AT}}": verified_at.astimezone(UTC).isoformat(timespec="seconds"),
        "{{VERIFIED_BY}}": verified_by_username,
    }


def generate_registration_document(
    registration: Registration,
    verified_by_username: str,
    template_path: Path,
) -> bytes:
    try:
        return _generate_registration_document(
            registration, verified_by_username, template_path
        )
    except DocumentTemplateError:
        raise
    except Exception as error:
        raise DocumentTemplateError(TEMPLATE_ERROR_MESSAGE) from error


def _generate_registration_document(
    registration: Registration,
    verified_by_username: str,
    template_path: Path,
) -> bytes:
    document = Document(template_path)
    replacements = replacement_values(registration, verified_by_username)
    matching_runs: dict[str, list[Run]] = {
        placeholder: [] for placeholder in replacements
    }
    paragraphs = list(iter_paragraphs(document))
    discovered_tokens = [
        token
        for paragraph in paragraphs
        for token in PLACEHOLDER_PATTERN.findall(paragraph.text)
    ]
    if Counter(discovered_tokens) != Counter(replacements.keys()):
        raise DocumentTemplateError(TEMPLATE_ERROR_MESSAGE)

    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.text in matching_runs:
                matching_runs[run.text].append(run)

    if any(len(runs) != 1 for runs in matching_runs.values()):
        raise DocumentTemplateError(TEMPLATE_ERROR_MESSAGE)

    for placeholder, value in replacements.items():
        matching_runs[placeholder][0].text = value

    scrub_core_properties(document)
    output = BytesIO()
    document.save(output)
    return scrub_extended_properties(output.getvalue())


def scrub_core_properties(document: DocumentObject) -> None:
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.keywords = ""
    properties.comments = ""
    for property_name in ("lastPrinted", "revision"):
        element = getattr(properties._element, property_name)
        if element is not None:
            properties._element.remove(element)


def scrub_extended_properties(content: bytes) -> bytes:
    source_buffer = BytesIO(content)
    output_buffer = BytesIO()
    with ZipFile(source_buffer) as source, ZipFile(
        output_buffer, "w", ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            entry_content = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                entry_content = scrub_app_properties(entry_content)
            target.writestr(entry, entry_content)
    return output_buffer.getvalue()


def scrub_app_properties(content: bytes) -> bytes:
    root = etree.fromstring(content)
    namespace = root.nsmap[None]
    for property_name in ("Company", "Manager"):
        element = root.find(f"{{{namespace}}}{property_name}")
        if element is not None:
            element.text = None
    application = root.find(f"{{{namespace}}}Application")
    if application is not None:
        application.text = "HSK Registration"
    app_version = root.find(f"{{{namespace}}}AppVersion")
    if app_version is not None:
        app_version.text = None
    return etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
