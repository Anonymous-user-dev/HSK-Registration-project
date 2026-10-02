from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLACK = "000000"
DARK_GRAY = "404040"
MID_GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
FONT_NAME = "Arial"


def set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: str = BLACK,
) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_GRAY) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_spacing(paragraph, *, before: float = 0, after: float = 0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0


def remove_paragraph_borders(properties) -> None:
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def add_section_heading(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    set_paragraph_spacing(paragraph, before=7, after=3)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=True)


def fill_field_cell(cell, label: str, value: str) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=80, start=120, bottom=90, end=120)
    label_paragraph = cell.paragraphs[0]
    set_paragraph_spacing(label_paragraph, after=2)
    label_run = label_paragraph.add_run(label)
    set_run_font(label_run, size=8, bold=True, color=MID_GRAY)

    value_paragraph = cell.add_paragraph()
    set_paragraph_spacing(value_paragraph)
    value_run = value_paragraph.add_run(value)
    set_run_font(value_run, size=10, color=DARK_GRAY)


def add_field_table(
    document: DocumentObject, rows: list[tuple[tuple[str, str], tuple[str, str]]]
) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row, fields in zip(table.rows, rows, strict=True):
        for cell, (label, value) in zip(row.cells, fields, strict=True):
            cell.width = Inches(3.35)
            fill_field_cell(cell, label, value)


def add_top_reference(document: DocumentObject) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for cell in table.rows[0].cells:
        cell.width = Inches(3.35)
        set_cell_shading(cell, PALE_BLUE)
    fill_field_cell(table.cell(0, 0), "Registration code", "{{REGISTRATION_CODE}}")
    fill_field_cell(table.cell(0, 1), "Record status", "VERIFIED")


def add_signature_table(document: DocumentObject) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for cell, label in zip(
        table.rows[0].cells,
        ("Candidate signature and date", "Staff signature and date"),
        strict=True,
    ):
        cell.width = Inches(3.35)
        set_cell_margins(cell, top=90, start=120, bottom=90, end=120)
        paragraph = cell.paragraphs[0]
        set_paragraph_spacing(paragraph, after=14)
        label_run = paragraph.add_run(label)
        set_run_font(label_run, size=8, bold=True, color=MID_GRAY)
        line = cell.add_paragraph("________________________________")
        set_paragraph_spacing(line)
        for run in line.runs:
            set_run_font(run, size=9, color=DARK_GRAY)


def add_notes_area(document: DocumentObject) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    cell = table.cell(0, 0)
    cell.width = Inches(6.7)
    set_cell_margins(cell, top=75, start=120, bottom=75, end=120)
    paragraph = cell.paragraphs[0]
    set_paragraph_spacing(paragraph, after=5)
    label = paragraph.add_run("Office notes")
    set_run_font(label, size=8, bold=True, color=MID_GRAY)
    notes = cell.add_paragraph("\n")
    set_paragraph_spacing(notes)


def configure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_after = Pt(4)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)

    title = document.styles["Title"]
    title.font.name = FONT_NAME
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    remove_paragraph_borders(title._element.get_or_add_pPr())

    heading = document.styles["Heading 1"]
    heading.font.name = FONT_NAME
    heading.font.size = Pt(10.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(BLACK)
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)


def build_template(output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    configure_styles(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title, after=3)
    remove_paragraph_borders(title._p.get_or_add_pPr())
    title.add_run("HSK Candidate Registration Record")

    introduction = document.add_paragraph()
    set_paragraph_spacing(introduction, after=7)
    intro_run = introduction.add_run(
        "Staff use this record after comparing the candidate's submitted details "
        "with the physical identification presented at registration."
    )
    set_run_font(intro_run, size=9, color=MID_GRAY)

    add_top_reference(document)

    add_section_heading(document, "Candidate identity")
    add_field_table(
        document,
        [
            (("First name", "{{FIRST_NAME}}"), ("Last name", "{{LAST_NAME}}")),
            (
                ("Date of birth", "{{DATE_OF_BIRTH}}"),
                ("Nationality", "{{NATIONALITY}}"),
            ),
            (
                ("Passport number", "{{PASSPORT_NUMBER}}"),
                ("HSK level", "{{HSK_LEVEL}}"),
            ),
        ],
    )

    add_section_heading(document, "Contact information")
    add_field_table(
        document,
        [
            (
                ("Phone number", "{{PHONE_NUMBER}}"),
                ("Parent phone number", "{{PARENT_PHONE_NUMBER}}"),
            )
        ],
    )

    add_section_heading(document, "Verification")
    add_field_table(
        document,
        [
            (
                ("Verified at", "{{VERIFIED_AT}}"),
                ("Verified by", "{{VERIFIED_BY}}"),
            )
        ],
    )

    add_section_heading(document, "Signatures")
    add_signature_table(document)
    add_notes_area(document)

    privacy = document.add_paragraph()
    set_paragraph_spacing(privacy, before=6)
    privacy_run = privacy.add_run(
        "Sensitive personal information. Handle only by authorized institute staff "
        "and store or dispose of this record according to institute policy."
    )
    set_run_font(privacy_run, size=7.5, color=MID_GRAY)

    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "HSK Candidate Registration Record"
    properties.subject = "Verified candidate registration record"
    properties.keywords = ""
    properties.comments = ""
    properties.created = datetime(2000, 1, 1)
    properties.modified = datetime(2000, 1, 1)
    properties.revision = 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the HSK DOCX template")
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    build_template(arguments.output)


if __name__ == "__main__":
    main()
