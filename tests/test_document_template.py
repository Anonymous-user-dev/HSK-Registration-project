from __future__ import annotations

import re
from collections.abc import Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.config import Settings

REQUIRED_PLACEHOLDERS = {
    "{{REGISTRATION_CODE}}",
    "{{FIRST_NAME}}",
    "{{LAST_NAME}}",
    "{{DATE_OF_BIRTH}}",
    "{{NATIONALITY}}",
    "{{PASSPORT_NUMBER}}",
    "{{PHONE_NUMBER}}",
    "{{PARENT_PHONE_NUMBER}}",
    "{{HSK_LEVEL}}",
    "{{VERIFIED_AT}}",
    "{{VERIFIED_BY}}",
}
PLACEHOLDER_PATTERN = re.compile(r"{{[^{}]+}}")


def iter_paragraphs(parent: DocumentObject | _Cell) -> Iterator[Paragraph]:
    yield from parent.paragraphs
    for table in parent.tables:
        yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from iter_paragraphs(cell)


def test_bundled_template_has_exact_isolated_placeholders() -> None:
    document = Document(Settings.from_env({}).document_template_path)
    paragraphs = list(iter_paragraphs(document))
    placeholder_runs = [
        run.text
        for paragraph in paragraphs
        for run in paragraph.runs
        if PLACEHOLDER_PATTERN.fullmatch(run.text)
    ]
    all_tokens = [
        token
        for paragraph in paragraphs
        for token in PLACEHOLDER_PATTERN.findall(paragraph.text)
    ]

    assert set(all_tokens) == REQUIRED_PLACEHOLDERS
    assert len(all_tokens) == len(REQUIRED_PLACEHOLDERS)
    assert sorted(placeholder_runs) == sorted(REQUIRED_PLACEHOLDERS)
