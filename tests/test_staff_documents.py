from __future__ import annotations

import re
from collections.abc import Iterator
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session, sessionmaker

from app.config import Settings
from app.main import create_app
from app.models import (
    AuditAction,
    AuditLog,
    Registration,
    RegistrationStatus,
    StaffUser,
)
from app.schemas import RegistrationCreate
from app.security import hash_password
from app.services.registration_service import create_registration
from app.services.verification_service import reject_registration, verify_registration

CSRF_PATTERN = re.compile(r'name="csrf_token" value="([^"]+)"')
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def extract_csrf(response) -> str:
    match = CSRF_PATTERN.search(response.text)
    assert match is not None
    return match.group(1)


def login(client: TestClient, username: str = "registrar") -> None:
    token = extract_csrf(client.get("/staff/login"))
    response = client.post(
        "/staff/login",
        data={
            "username": username,
            "password": "correct horse battery",
            "csrf_token": token,
        },
        follow_redirects=False,
    )
    assert response.status_code == 303


@pytest.fixture
def document_app(tmp_path: Path) -> FastAPI:
    app = create_app(
        Settings(
            environment="test",
            database_url=f"sqlite:///{tmp_path / 'documents.db'}",
            session_secret="test-session-secret-that-is-not-for-production",
            secure_cookies=False,
        )
    )
    with app.state.session_factory() as db:
        staff = StaffUser(
            username="registrar",
            password_hash=hash_password("correct horse battery"),
            role="REGISTRATION_STAFF",
            is_active=True,
        )
        viewer = StaffUser(
            username="viewer",
            password_hash=hash_password("correct horse battery"),
            role="VIEW_ONLY",
            is_active=True,
        )
        inactive = StaffUser(
            username="inactive",
            password_hash=hash_password("correct horse battery"),
            role="REGISTRATION_STAFF",
            is_active=True,
        )
        db.add_all([staff, viewer, inactive])
        db.commit()
        db.refresh(staff)
        verified = create_registration(
            db,
            RegistrationCreate.model_validate(
                {
                    "first_name": "Mei",
                    "last_name": "Lin",
                    "date_of_birth": "2002-05-14",
                    "nationality": "Tajikistani",
                    "passport_number": "FAKE12345",
                    "phone_number": "+992 900 000 001",
                    "parent_phone_number": "+992 900 000 002",
                    "hsk_level": "HSK3",
                }
            ),
        )
        verify_registration(db, verified.id, staff.id, verified.version)
        pending = create_registration(
            db,
            RegistrationCreate.model_validate(
                {
                    "first_name": "Arman",
                    "last_name": "Karim",
                    "date_of_birth": "2001-04-12",
                    "nationality": "Tajikistani",
                    "passport_number": "FAKE67890",
                    "phone_number": "+992 900 000 003",
                    "parent_phone_number": "+992 900 000 004",
                    "hsk_level": "HSK2",
                }
            ),
        )
        rejected = create_registration(
            db,
            RegistrationCreate.model_validate(
                {
                    "first_name": "Sara",
                    "last_name": "Nazar",
                    "date_of_birth": "2003-08-20",
                    "nationality": "Tajikistani",
                    "passport_number": "FAKE24680",
                    "phone_number": "+992 900 000 005",
                    "parent_phone_number": "+992 900 000 006",
                    "hsk_level": "HSK4",
                }
            ),
        )
        reject_registration(db, rejected.id, staff.id, rejected.version)
        app.state.test_registration_code = verified.registration_code
        app.state.pending_registration_code = pending.registration_code
        app.state.rejected_registration_code = rejected.registration_code
    return app


@pytest.fixture
def document_client(document_app: FastAPI) -> Iterator[TestClient]:
    with TestClient(document_app) as client:
        login(client)
        yield client


def test_verified_detail_shows_csrf_protected_document_action(
    document_client: TestClient, document_app: FastAPI
) -> None:
    code = document_app.state.test_registration_code

    response = document_client.get(f"/staff/registrations/{code}")

    assert response.status_code == 200
    assert f'action="/staff/registrations/{code}/document"' in response.text
    assert 'name="csrf_token"' in response.text
    assert "Generate Word document" in response.text


def test_verified_document_download_returns_safe_no_store_attachment(
    document_client: TestClient, document_app: FastAPI
) -> None:
    code = document_app.state.test_registration_code
    detail = document_client.get(f"/staff/registrations/{code}")

    response = document_client.post(
        f"/staff/registrations/{code}/document",
        data={"csrf_token": extract_csrf(detail)},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_CONTENT_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="hsk_registration_{code}.docx"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "Mei" not in response.headers["content-disposition"]
    assert "FAKE12345" not in response.headers["content-disposition"]
    generated = Document(BytesIO(response.content))
    assert code in "\n".join(
        paragraph.text
        for table in generated.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    with document_app.state.session_factory() as db:
        assert (
            db.scalar(
                select(func.count(AuditLog.id)).where(
                    AuditLog.action == AuditAction.DOCUMENT_GENERATED
                )
            )
            == 1
        )


def document_post(client: TestClient, code: str, token: str | None = None):
    if token is None:
        token = extract_csrf(client.get("/staff"))
    return client.post(
        f"/staff/registrations/{code}/document",
        data={"csrf_token": token},
        follow_redirects=False,
    )


def document_audits(app: FastAPI) -> list[AuditLog]:
    with app.state.session_factory() as db:
        return list(
            db.scalars(
                select(AuditLog)
                .where(AuditLog.action == AuditAction.DOCUMENT_GENERATED)
                .order_by(AuditLog.id)
            )
        )


def test_anonymous_document_request_redirects_before_record_lookup(
    document_app: FastAPI,
) -> None:
    code = document_app.state.test_registration_code
    with TestClient(document_app) as client:
        response = document_post(client, code, "not-a-token")

    assert response.status_code == 303
    assert response.headers["location"] == "/staff/login"
    assert "FAKE12345" not in response.text
    assert document_audits(document_app) == []


def test_inactive_staff_cannot_generate_document(document_app: FastAPI) -> None:
    code = document_app.state.test_registration_code
    with TestClient(document_app) as client:
        login(client, "inactive")
        token = extract_csrf(client.get("/staff"))
        with document_app.state.session_factory() as db:
            user = db.scalar(select(StaffUser).where(StaffUser.username == "inactive"))
            assert user is not None
            user.is_active = False
            db.commit()

        response = document_post(client, code, token)

    assert response.status_code == 303
    assert response.headers["location"] == "/staff/login"
    assert "FAKE12345" not in response.text
    assert document_audits(document_app) == []


def test_view_only_staff_cannot_generate_or_receive_details(
    document_app: FastAPI,
) -> None:
    code = document_app.state.test_registration_code
    with TestClient(document_app) as client:
        login(client, "viewer")

        response = document_post(client, code)

    assert response.status_code == 403
    assert "FAKE12345" not in response.text
    assert document_audits(document_app) == []


@pytest.mark.parametrize("token", ["", "wrong-token"])
def test_invalid_csrf_blocks_document_generation(
    document_client: TestClient,
    document_app: FastAPI,
    token: str,
) -> None:
    response = document_post(
        document_client, document_app.state.test_registration_code, token
    )

    assert response.status_code == 403
    assert document_audits(document_app) == []


def test_unknown_registration_cannot_generate_document(
    document_client: TestClient, document_app: FastAPI
) -> None:
    response = document_post(document_client, "HSK-ZZZZZZ")

    assert response.status_code == 404
    assert "FAKE12345" not in response.text
    assert document_audits(document_app) == []


@pytest.mark.parametrize(
    "code_attribute",
    ["pending_registration_code", "rejected_registration_code"],
)
def test_unverified_registration_cannot_generate_document(
    document_client: TestClient,
    document_app: FastAPI,
    code_attribute: str,
) -> None:
    code = getattr(document_app.state, code_attribute)

    detail = document_client.get(f"/staff/registrations/{code}")
    response = document_post(document_client, code)

    assert response.status_code == 409
    assert DOCX_CONTENT_TYPE not in response.headers.get("content-type", "")
    assert "Generate Word document" not in detail.text
    assert document_audits(document_app) == []


def test_missing_original_verifier_fails_closed(
    document_client: TestClient, document_app: FastAPI
) -> None:
    code = document_app.state.test_registration_code
    with document_app.state.session_factory() as db:
        registration = db.scalar(
            select(Registration).where(Registration.registration_code == code)
        )
        assert registration is not None
        registration.verified_by = 99999
        db.commit()

    response = document_post(document_client, code)

    assert response.status_code == 503
    assert "FAKE12345" not in response.text
    assert document_audits(document_app) == []


@pytest.mark.parametrize("template_content", [None, b"not a Word document"])
def test_unavailable_template_returns_generic_failure_without_audit(
    tmp_path: Path,
    document_client: TestClient,
    document_app: FastAPI,
    template_content: bytes | None,
) -> None:
    template_path = tmp_path / "unavailable-template.docx"
    if template_content is not None:
        template_path.write_bytes(template_content)
    document_app.state.settings = replace(
        document_app.state.settings, document_template_path=template_path
    )

    response = document_post(document_client, document_app.state.test_registration_code)

    assert response.status_code == 503
    assert "FAKE12345" not in response.text
    assert str(template_path) not in response.text
    assert DOCX_CONTENT_TYPE not in response.headers.get("content-type", "")
    assert document_audits(document_app) == []


def test_repeated_downloads_create_separate_redacted_audits(
    document_client: TestClient, document_app: FastAPI
) -> None:
    code = document_app.state.test_registration_code

    first = document_post(document_client, code)
    second = document_post(document_client, code)

    assert first.status_code == 200
    assert second.status_code == 200
    audits = document_audits(document_app)
    assert len(audits) == 2
    assert audits[0].id != audits[1].id
    assert [audit.changed_fields for audit in audits] == [[], []]
    with document_app.state.session_factory() as db:
        registration = db.scalar(
            select(Registration).where(Registration.registration_code == code)
        )
        assert registration is not None
        assert registration.status is RegistrationStatus.VERIFIED
        assert registration.version == 2


def test_audit_commit_failure_prevents_document_response(
    document_client: TestClient, document_app: FastAPI
) -> None:
    code = document_app.state.test_registration_code
    working_factory = document_app.state.session_factory

    class FailingAuditSession(Session):
        def commit(self) -> None:
            if any(
                isinstance(item, AuditLog)
                and item.action is AuditAction.DOCUMENT_GENERATED
                for item in self.new
            ):
                raise SQLAlchemyError("simulated audit failure")
            super().commit()

    document_app.state.session_factory = sessionmaker(
        bind=document_app.state.engine,
        class_=FailingAuditSession,
        expire_on_commit=False,
    )

    response = document_post(document_client, code)

    assert response.status_code == 500
    assert DOCX_CONTENT_TYPE not in response.headers.get("content-type", "")
    with working_factory() as db:
        assert (
            db.scalar(
                select(func.count(AuditLog.id)).where(
                    AuditLog.action == AuditAction.DOCUMENT_GENERATED
                )
            )
            == 0
        )
