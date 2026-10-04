from __future__ import annotations

import re
from collections.abc import Callable, Iterator
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

from app.config import Settings
from app.models import HSKLevel, Registration, RegistrationStatus
from app.services.document_service import (
    DocumentTemplateError,
    generate_registration_document,
)


def verified_registration() -> Registration:
    return Registration(
        id=1,
        registration_code="HSK-FAKE01",
        first_name="Mei",
        last_name="Lin",
        date_of_birth=date(2002, 5, 14),
        nationality="Tajikistani",
        passport_number="FAKE12345",
        phone_number="+992 900 000 001",
        parent_phone_number="+992 900 000 002",
        hsk_level=HSKLevel.HSK3,
        status=RegistrationStatus.VERIFIED,
        created_at=datetime(2026, 10, 1, 12, 0, tzinfo=UTC),
        verified_at=datetime(2026, 10, 2, 17, 30, tzinfo=UTC),
        verified_by=1,
        version=2,
    )


def iter_paragraphs(parent: DocumentObject | _Cell) -> Iterator[Paragraph]:
    yield from parent.paragraphs
    for table in parent.tables:
        yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table: Table) -> Iterator[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            yield from iter_paragraphs(cell)


def document_text(document: DocumentObject) -> str:
    return "\n".join(paragraph.text for paragraph in iter_paragraphs(document))


def find_placeholder_run(document: DocumentObject, placeholder: str):
    return next(
        run
        for paragraph in iter_paragraphs(document)
        for run in paragraph.runs
        if run.text == placeholder
    )


def modified_template(
    tmp_path: Path,
    modify: Callable[[DocumentObject], None],
) -> Path:
    document = Document(Settings.from_env({}).document_template_path)
    modify(document)
    template_path = tmp_path / "modified-template.docx"
    document.save(template_path)
    return template_path


def generate_from_template(template_path: Path) -> bytes:
    return generate_registration_document(
        verified_registration(), "registrar", template_path
    )


def test_generate_registration_document_returns_valid_docx_with_fake_values(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.chdir(tmp_path)

    content = generate_registration_document(
        verified_registration(),
        "registrar",
        Settings.from_env({}).document_template_path,
    )

    document = Document(BytesIO(content))
    text = document_text(document)
    for expected in (
        "HSK-FAKE01",
        "Mei",
        "Lin",
        "2002-05-14",
        "Tajikistani",
        "FAKE12345",
        "+992 900 000 001",
        "+992 900 000 002",
        "HSK3",
        "2026-10-02T17:30:00+00:00",
        "registrar",
    ):
        assert expected in text
    assert re.search(r"{{[^{}]+}}", text) is None
    assert "Sensitive personal information" in text
    registration_label = next(
        run
        for paragraph in iter_paragraphs(document)
        for run in paragraph.runs
        if run.text == "Registration code"
    )
    assert registration_label.bold is True
    assert list(tmp_path.glob("*.docx")) == []


def test_missing_template_is_rejected_with_domain_error(tmp_path: Path) -> None:
    with pytest.raises(
        DocumentTemplateError, match="Document template could not be processed"
    ):
        generate_from_template(tmp_path / "missing.docx")


def test_malformed_template_is_rejected_with_domain_error(tmp_path: Path) -> None:
    template_path = tmp_path / "malformed.docx"
    template_path.write_bytes(b"not a Word document")

    with pytest.raises(
        DocumentTemplateError, match="Document template could not be processed"
    ):
        generate_from_template(template_path)


def test_template_missing_required_placeholder_is_rejected(tmp_path: Path) -> None:
    def remove_placeholder(document: DocumentObject) -> None:
        find_placeholder_run(document, "{{FIRST_NAME}}").text = ""

    template_path = modified_template(tmp_path, remove_placeholder)

    with pytest.raises(DocumentTemplateError):
        generate_from_template(template_path)


def test_template_with_duplicate_placeholder_is_rejected(tmp_path: Path) -> None:
    def duplicate_placeholder(document: DocumentObject) -> None:
        document.add_paragraph().add_run("{{FIRST_NAME}}")

    template_path = modified_template(tmp_path, duplicate_placeholder)

    with pytest.raises(DocumentTemplateError):
        generate_from_template(template_path)


def test_template_with_unknown_placeholder_is_rejected(tmp_path: Path) -> None:
    def add_unknown_placeholder(document: DocumentObject) -> None:
        document.add_paragraph().add_run("{{UNEXPECTED_VALUE}}")

    template_path = modified_template(tmp_path, add_unknown_placeholder)

    with pytest.raises(DocumentTemplateError):
        generate_from_template(template_path)


def test_template_with_placeholder_split_across_runs_is_rejected(
    tmp_path: Path,
) -> None:
    def split_placeholder(document: DocumentObject) -> None:
        run = find_placeholder_run(document, "{{REGISTRATION_CODE}}")
        run.text = "{{REGISTRATION_"
        run._parent.add_run("CODE}}")

    template_path = modified_template(tmp_path, split_placeholder)

    with pytest.raises(DocumentTemplateError):
        generate_from_template(template_path)


def test_generated_document_scrubs_personal_build_metadata(tmp_path: Path) -> None:
    document = Document(Settings.from_env({}).document_template_path)
    properties = document.core_properties
    properties.author = "Fake Developer"
    properties.last_modified_by = "Fake Machine"
    properties.last_printed = datetime(2026, 10, 2, 8, 0)
    properties.revision = 99
    seeded_path = tmp_path / "seeded.docx"
    document.save(seeded_path)

    rewritten_path = tmp_path / "seeded-extended.docx"
    with ZipFile(seeded_path) as source, ZipFile(
        rewritten_path, "w", ZIP_DEFLATED
    ) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                root = etree.fromstring(content)
                namespace = root.nsmap[None]
                company = root.find(f"{{{namespace}}}Company")
                assert company is not None
                company.text = "Fake Company"
                manager = root.find(f"{{{namespace}}}Manager")
                if manager is None:
                    manager = etree.SubElement(root, f"{{{namespace}}}Manager")
                manager.text = "Fake Manager"
                content = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            target.writestr(entry, content)

    result = generate_from_template(rewritten_path)

    with ZipFile(BytesIO(result)) as generated:
        metadata = b"".join(
            generated.read(name)
            for name in ("docProps/core.xml", "docProps/app.xml")
        )
    for secret in (
        b"Fake Developer",
        b"Fake Machine",
        b"Fake Company",
        b"Fake Manager",
        b"2026-10-02",
        b">99<",
    ):
        assert secret not in metadata
